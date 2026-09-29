"""
make_study_sheet.py — create coin_study.xlsx: one row per photo.

  Yellow  (you fill)      : real coins / real non-coin round objects in the photo
  Grey    (machine part)  : filled from the program's output + circle-by-circle review
  Blue    (formulas)      : error, TP, precision, recall — recalculate automatically
"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
if (ROOT / "coin_study.xlsx").exists():
    raise SystemExit("coin_study.xlsx already exists and holds the filled-in study; delete or rename it first.")
stats = json.loads((ROOT / "report" / "figures" / "stats.json").read_text(encoding="utf-8"))
rows = sorted(stats["rows"], key=lambda r: r["id"])  # IDs and file names only

FONT = "Arial"
YELLOW = PatternFill("solid", fgColor="FFF4C2")
GREY = PatternFill("solid", fgColor="E7EAEE")
BLUE = PatternFill("solid", fgColor="EAF2FC")
HEAD_HUMAN = PatternFill("solid", fgColor="E0B400")
HEAD_MACHINE = PatternFill("solid", fgColor="5F6B7A")
HEAD_CALC = PatternFill("solid", fgColor="1565C0")
HEAD_INFO = PatternFill("solid", fgColor="0B3C74")
thin = Side(style="thin", color="C9D2DE")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


wb = Workbook()

# ── Sheet 1: per-photo table ─────────────────────────────────────────
ws = wb.active
ws.title = "Photos"
ws.sheet_view.rightToLeft = True

ws["A1"] = "טבלת מחקר: מה אני ספרתי מול מה המכונה זיהתה"
ws["A1"].font = f(True, "0B3C74", 14)
ws["A2"] = ("צהוב = את/ה ממלא/ת (ספירה אמיתית מהתמונה המקורית).   אפור = החלק של המכונה (ימולא מהפלט של התוכנית).   "
            "כחול = נוסחאות, לא לגעת.   דוגמה למילוי בגיליון 'How to fill'.")
ws["A2"].font = f(color="5F6B7A", size=9)

# Group header row 4, column header row 5, data from row 6
groups = [(1, 4, "פרטי התמונה", HEAD_INFO), (5, 7, "אני (ספירה אמיתית)", HEAD_HUMAN),
          (8, 11, "המכונה", HEAD_MACHINE), (12, 15, "השוואה (נוסחאות)", HEAD_CALC)]
for c1, c2, text, fill in groups:
    ws.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c2)
    cell = ws.cell(row=4, column=c1, value=text)
    cell.font = f(True, "FFFFFF", 11)
    cell.fill = fill
    cell.alignment = CENTER

headers = [
    ("מזהה", 7), ("קובץ תמונה", 34), ("קטגוריה", 10), ("תמונת תוצאה", 14),
    ("מטבעות אמיתיים בתמונה", 13), ("עצמים עגולים שאינם מטבע בתמונה", 15), ("הערות", 24),
    ("מטבעות שהמכונה ספרה", 13), ("זיהוי שגוי: לא-מטבע שנספר כמטבע (FP)", 17),
    ("מטבעות שהמכונה פספסה (FN)", 14), ("סוג הזיהויים השגויים", 22),
    ("טעות בספירה (מכונה − אמת)", 13), ("זיהויים נכונים (TP)", 11), ("Precision", 11), ("Recall", 11),
]
fills = [None] * 4 + [YELLOW] * 3 + [GREY] * 4 + [BLUE] * 4
for i, (h, w) in enumerate(headers, 1):
    c = ws.cell(row=5, column=i, value=h)
    c.font = f(True, size=10)
    c.alignment = CENTER
    c.border = BORDER
    c.fill = fills[i - 1] or PatternFill("solid", fgColor="D6E4F5")
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[5].height = 48

comments = {
    5: "כמה מטבעות יש באמת בתמונה, לפי התמונה המקורית (עמודה B).",
    6: "כפתורים או כל עצם עגול אחר שאינו מטבע ונמצא על המשטח.",
    9: "עיגול ירוק על משהו שאינו מטבע: קפל, כפתור, צל, או מטבע שנספר פעמיים.",
    10: "מטבע אמיתי שאין עליו עיגול ירוק בתמונת התוצאה.",
    11: "למשל: קפל ×2, כפתור ×1, כפול ×1",
    13: "TP = מטבעות שהמכונה ספרה − זיהויים שגויים",
    14: "Precision = TP / מטבעות שהמכונה ספרה: כמה מהסימונים נכונים",
    15: "Recall = TP / מטבעות אמיתיים: כמה מהמטבעות נמצאו",
}
for col, text in comments.items():
    ws.cell(row=5, column=col).comment = Comment(text, "coin study")

FIRST = 6
for k, r in enumerate(rows):
    row = FIRST + k
    fname = f"WhatsApp Image 2026-09-01 at {r['short']}.jpeg"
    ws.cell(row=row, column=1, value=r["id"]).font = f(True, "0B3C74")
    c = ws.cell(row=row, column=2, value=fname)
    c.hyperlink = f"data/{fname}"
    c.font = f(color="1565C0")
    ws.cell(row=row, column=3, value=r["category"]).font = f()
    c = ws.cell(row=row, column=4, value="פתח")
    c.hyperlink = f"results/{Path(fname).stem}.jpg"
    c.font = f(color="1565C0")
    for col in range(5, 12):
        ws.cell(row=row, column=col).fill = YELLOW if col <= 7 else GREY
    # Formulas: blank until the inputs they need are filled
    E, H, I = f"E{row}", f"H{row}", f"I{row}"
    ws.cell(row=row, column=12, value=f'=IF(OR({E}="",{H}=""),"",{H}-{E})')
    ws.cell(row=row, column=13, value=f'=IF(OR({H}="",{I}=""),"",{H}-{I})')
    ws.cell(row=row, column=14, value=f'=IF(OR(M{row}="",{H}=""),"",IFERROR(M{row}/{H},""))')
    ws.cell(row=row, column=15, value=f'=IF(OR(M{row}="",{E}=""),"",IFERROR(M{row}/{E},""))')
    for col in range(1, 16):
        cell = ws.cell(row=row, column=col)
        cell.border = BORDER
        if col not in (2, 7, 11):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        if col >= 12:
            cell.fill = BLUE
            cell.font = f()
        elif col >= 5:
            cell.font = f(color="0000FF")
    ws.cell(row=row, column=12).number_format = '+0;-0;0'
    ws.cell(row=row, column=14).number_format = "0.0%"
    ws.cell(row=row, column=15).number_format = "0.0%"
LAST = FIRST + len(rows) - 1

# Totals row
tot = LAST + 1
ws.cell(row=tot, column=1, value="סה\"כ").font = f(True)
for col in (5, 6, 8, 9, 10):
    L = get_column_letter(col)
    ws.cell(row=tot, column=col, value=f'=IF(COUNT({L}{FIRST}:{L}{LAST})=0,"",SUM({L}{FIRST}:{L}{LAST}))')
ws.cell(row=tot, column=12, value=f'=IF(OR(E{tot}="",H{tot}=""),"",H{tot}-E{tot})')
ws.cell(row=tot, column=13, value=f'=IF(COUNT(M{FIRST}:M{LAST})=0,"",SUM(M{FIRST}:M{LAST}))')
ws.cell(row=tot, column=14, value=f'=IF(OR(M{tot}="",H{tot}=""),"",IFERROR(M{tot}/H{tot},""))')
ws.cell(row=tot, column=15, value=f'=IF(OR(M{tot}="",E{tot}=""),"",IFERROR(M{tot}/E{tot},""))')
for col in range(1, 16):
    cell = ws.cell(row=tot, column=col)
    cell.border = Border(top=Side(style="medium", color="0B3C74"), bottom=thin, left=thin, right=thin)
    cell.fill = PatternFill("solid", fgColor="D6E4F5")
    cell.font = f(True)
    cell.alignment = Alignment(horizontal="center", vertical="center")
ws.cell(row=tot, column=12).number_format = '+0;-0;0'
ws.cell(row=tot, column=14).number_format = "0.0%"
ws.cell(row=tot, column=15).number_format = "0.0%"
ws.freeze_panes = "E6"

# ── Sheet 2: summary by category ────────────────────────────────────
sm = wb.create_sheet("Summary")
sm.sheet_view.rightToLeft = True
sm["A1"] = "סיכום לפי קטגוריה (מתעדכן אוטומטית מגיליון Photos)"
sm["A1"].font = f(True, "0B3C74", 14)
cols = ["קטגוריה", "תמונות", "מטבעות אמיתיים", "המכונה ספרה", "זיהויים נכונים (TP)", "זיהויים שגויים (FP)",
        "פספוסים (FN)", "Precision", "Recall", "תמונות עם ספירה מדויקת"]
for i, h in enumerate(cols, 1):
    c = sm.cell(row=3, column=i, value=h)
    c.font = f(True, "FFFFFF")
    c.fill = HEAD_CALC
    c.alignment = CENTER
    c.border = BORDER
    sm.column_dimensions[get_column_letter(i)].width = 15
sm.row_dimensions[3].height = 34
rng = lambda L: f"Photos!${L}${FIRST}:${L}${LAST}"  # noqa: E731
for k, cat in enumerate(["clean", "noisy", "damaged", "הכל"]):
    r = 4 + k
    sm.cell(row=r, column=1, value=cat)
    crit = '"*"' if cat == "הכל" else f"$A{r}"
    sm.cell(row=r, column=2, value=f"=COUNTIFS({rng('C')},{crit})")
    for col, L in ((3, "E"), (4, "H"), (5, "M"), (6, "I"), (7, "J")):
        sm.cell(row=r, column=col, value=f"=SUMIFS({rng(L)},{rng('C')},{crit})")
    sm.cell(row=r, column=8, value=f'=IF(D{r}=0,"",E{r}/D{r})')
    sm.cell(row=r, column=9, value=f'=IF(C{r}=0,"",E{r}/C{r})')
    sm.cell(row=r, column=10, value=f'=COUNTIFS({rng("C")},{crit},{rng("L")},0)')
    for col in range(1, 11):
        cell = sm.cell(row=r, column=col)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
        cell.font = f(bold=(cat == "הכל"))
        if cat == "הכל":
            cell.fill = PatternFill("solid", fgColor="D6E4F5")
    sm.cell(row=r, column=8).number_format = "0.0%"
    sm.cell(row=r, column=9).number_format = "0.0%"
sm["A10"] = "Precision: מתוך כל הסימונים של המכונה, כמה היו מטבעות אמיתיים.   Recall: מתוך כל המטבעות האמיתיים, כמה המכונה מצאה."
sm["A10"].font = f(color="5F6B7A", size=9)

# ── Sheet 3: how to fill + example ──────────────────────────────────
hw = wb.create_sheet("How to fill")
hw.sheet_view.rightToLeft = True
lines = [
    ("איך ממלאים", True),
    ("1. בגיליון Photos, לכל שורה: לוחצים על שם הקובץ בעמודה B כדי לפתוח את התמונה המקורית.", False),
    ("2. ממלאים בעמודות הצהובות: כמה מטבעות יש באמת, וכמה עצמים עגולים שאינם מטבע (כפתורים וכו').", False),
    ("3. את העמודות האפורות (המכונה) אני ממלא: כמה מטבעות ספרה, ובדיקה של כל עיגול בתמונת התוצאה.", False),
    ("4. העמודות הכחולות וגיליון Summary מחושבים אוטומטית.", False),
    ("", False),
    ("דוגמה לשורה ממולאת (לא חלק מהנתונים)", True),
]
for i, (t, b) in enumerate(lines, 1):
    hw.cell(row=i, column=1, value=t).font = f(bold=b, size=11 if b else 10, color="0B3C74" if b else "000000")
ex_head = ["מזהה", "מטבעות אמיתיים", "עצמים שאינם מטבע", "הערות", "המכונה ספרה", "FP", "FN", "סוג השגיאות", "טעות", "TP"]
ex_vals = ["P20", 14, 4, "4 כפתורים בין המטבעות", 17, 4, 1, "כפתור ×1, קפל ×2, קצה ×1", "=E9-B9", "=E9-F9"]
for i, (h, v) in enumerate(zip(ex_head, ex_vals), 1):
    a = hw.cell(row=8, column=i, value=h)
    a.font = f(True)
    a.fill = PatternFill("solid", fgColor="D6E4F5")
    b = hw.cell(row=9, column=i, value=v)
    b.font = f(italic=True, color="5F6B7A")
    for c in (a, b):
        c.border = BORDER
        c.alignment = CENTER
    hw.column_dimensions[get_column_letter(i)].width = 16
hw["A11"] = "המספרים בדוגמה להמחשה בלבד."
hw["A11"].font = f(color="5F6B7A", size=9)
hw.column_dimensions["A"].width = 16

out = ROOT / "coin_study.xlsx"
wb.save(out)
print("wrote", out)
